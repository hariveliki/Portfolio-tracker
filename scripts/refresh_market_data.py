#!/usr/bin/env python3
"""Refresh market data table in the workbook from IBKR (with optional yfinance fallback)."""

from __future__ import annotations

import argparse
import os
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any

import pandas as pd
import yaml
from ib_insync import IB, Stock, util
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.worksheet.table import Table, TableStyleInfo


DEFAULT_CONFIG = "config/portfolio.yaml"
MARKET_HEADERS = ["Date", "Ticker", "Open", "High", "Low", "Close", "Adj_Close", "Volume"]

HEADER_FILL = PatternFill(start_color="1C2541", end_color="1C2541", fill_type="solid")
HEADER_FONT = Font(name="Segoe UI", color="FFFFFF", bold=True, size=10)
TABLE_STYLE = TableStyleInfo(
    name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
    showRowStripes=True, showColumnStripes=False,
)


@dataclass
class RefreshConfig:
    workbook: Path
    config_path: Path
    ib_host: str
    ib_port: int
    ib_client_id: int
    yfinance_enabled: bool
    yfinance_default: bool
    yfinance_per_ticker: dict[str, bool]


@dataclass
class ContractHint:
    exchange: str
    currency: str


def load_yaml_config(config_path: Path) -> dict[str, Any]:
    if not config_path.exists():
        return {}
    with config_path.open("r", encoding="utf-8") as handle:
        return yaml.safe_load(handle) or {}


def parse_args() -> RefreshConfig:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook")
    parser.add_argument("--config", default=DEFAULT_CONFIG)
    parser.add_argument("--ib-host")
    parser.add_argument("--ib-port", type=int)
    parser.add_argument("--ib-client-id", type=int)
    args = parser.parse_args()

    config_path = Path(args.config)
    raw_cfg = load_yaml_config(config_path)
    ib_cfg = raw_cfg.get("ib", {})
    path_cfg = raw_cfg.get("paths", {})
    fallback_cfg = raw_cfg.get("yfinance_fallback", {})

    return RefreshConfig(
        workbook=Path(args.workbook or path_cfg.get("workbook_output", "output/portfolio_workbook.xlsx")),
        config_path=config_path,
        ib_host=args.ib_host or os.getenv("IB_HOST", str(ib_cfg.get("host", "127.0.0.1"))),
        ib_port=args.ib_port or int(os.getenv("IB_PORT", str(ib_cfg.get("port", 7497)))),
        ib_client_id=args.ib_client_id or int(os.getenv("IB_CLIENT_ID", str(ib_cfg.get("client_id", 15)))),
        yfinance_enabled=bool(fallback_cfg.get("enabled", True)),
        yfinance_default=bool(fallback_cfg.get("default", True)),
        yfinance_per_ticker={k.upper(): bool(v) for k, v in fallback_cfg.get("tickers", {}).items()},
    )


def read_table_dataframe(wb: Workbook, table_name: str) -> tuple[pd.DataFrame, Any, str]:
    for ws in wb.worksheets:
        if table_name in ws.tables:
            table = ws.tables[table_name]
            min_col, min_row, max_col, max_row = range_boundaries(table.ref)
            values = list(
                ws.iter_rows(
                    min_row=min_row, max_row=max_row,
                    min_col=min_col, max_col=max_col,
                    values_only=True,
                )
            )
            if not values:
                return pd.DataFrame(), ws, table_name
            header = values[0]
            rows = [r for r in values[1:] if any(v is not None for v in r)]
            return pd.DataFrame(rows, columns=header), ws, table_name
    raise ValueError(f"Table '{table_name}' not found in workbook")


def extract_tickers_start_and_hints(df: pd.DataFrame) -> tuple[list[str], date, dict[str, ContractHint]]:
    if "Ticker" not in df.columns:
        raise ValueError("tbl_TradeLog must have a 'Ticker' column")

    def normalize_ticker(value: Any) -> str:
        if pd.isna(value):
            return ""
        text = str(value).strip().upper()
        return "" if text in {"", "NAN"} or text.startswith("=") else text

    work = df.copy()
    work["Ticker"] = work["Ticker"].apply(normalize_ticker)
    work = work[work["Ticker"].ne("")]
    tickers = sorted(set(work["Ticker"]))

    start_candidates = []
    for col in ("Date",):
        if col in work.columns:
            parsed = pd.to_datetime(work[col], errors="coerce").dropna()
            if not parsed.empty:
                start_candidates.append(parsed.min().date())
    start_date = min(start_candidates) if start_candidates else date.today().replace(year=date.today().year - 1)

    hints: dict[str, ContractHint] = {}
    if "exchange" in work.columns or "currency" in work.columns:
        if "Date" in work.columns:
            work["Date"] = pd.to_datetime(work["Date"], errors="coerce")
            work = work.sort_values("Date")
        for _, row in work.iterrows():
            ticker = row["Ticker"]
            exchange = str(row.get("exchange") or "").strip().upper()
            currency = str(row.get("currency") or "").strip().upper()
            if not exchange and not currency:
                continue
            hints[ticker] = ContractHint(exchange=exchange, currency=currency)

    return tickers, start_date, hints


def _candidate_contracts(ticker: str, hint: ContractHint | None) -> list[Stock]:
    candidates: list[tuple[str, str]] = []
    if hint:
        if hint.exchange and hint.currency:
            candidates.append((hint.exchange, hint.currency))
        elif hint.exchange:
            candidates.append((hint.exchange, "USD"))
        elif hint.currency:
            candidates.append(("SMART", hint.currency))
        if hint.currency:
            candidates.append(("SMART", hint.currency))
    candidates.append(("SMART", "USD"))

    dedup: list[tuple[str, str]] = []
    seen: set[tuple[str, str]] = set()
    for item in candidates:
        if item in seen:
            continue
        seen.add(item)
        dedup.append(item)
    return [Stock(ticker, exchange, currency) for exchange, currency in dedup]


def fetch_ib_series(ib: IB, ticker: str, start_date: date, hint: ContractHint | None = None) -> pd.DataFrame:
    duration_days = max((date.today() - start_date).days + 5, 10)
    if duration_days > 365:
        duration_str = f"{(duration_days + 364) // 365} Y"
    else:
        duration_str = f"{duration_days} D"
    for contract in _candidate_contracts(ticker, hint):
        try:
            qualified = ib.qualifyContracts(contract)
            if not qualified:
                continue
            bars = ib.reqHistoricalData(
                contract,
                endDateTime="",
                durationStr=duration_str,
                barSizeSetting="1 day",
                whatToShow="ADJUSTED_LAST",
                useRTH=True,
                formatDate=1,
            )
        except Exception:
            continue

        df = util.df(bars)
        if df is None or df.empty:
            continue
        df = df[["date", "open", "high", "low", "close", "volume"]].copy()
        df.columns = ["Date", "Open", "High", "Low", "Close", "Volume"]
        df["Date"] = pd.to_datetime(df["Date"]).dt.date
        df = df[df["Date"] >= start_date]
        df["Adj_Close"] = df["Close"]
        df["Ticker"] = ticker
        if not df.empty:
            return df[MARKET_HEADERS]
    return pd.DataFrame()


def fetch_yf_series(ticker: str, start_date: date) -> pd.DataFrame:
    import yfinance as yf

    hist = yf.download(
        ticker,
        start=start_date.isoformat(),
        end=date.today().isoformat(),
        progress=False,
        auto_adjust=True,
    )
    if hist.empty:
        return pd.DataFrame()
    if isinstance(hist.columns, pd.MultiIndex):
        hist.columns = hist.columns.droplevel("Ticker")
    df = hist.reset_index()
    df = df[["Date", "Open", "High", "Low", "Close", "Volume"]].copy()
    df["Date"] = pd.to_datetime(df["Date"]).dt.date
    df["Adj_Close"] = df["Close"]
    df["Ticker"] = ticker
    return df[MARKET_HEADERS]


def should_use_yf(cfg: RefreshConfig, ticker: str) -> bool:
    if not cfg.yfinance_enabled:
        return False
    return cfg.yfinance_per_ticker.get(ticker.upper(), cfg.yfinance_default)


def write_market_table(wb: Workbook, market_df: pd.DataFrame, table_name: str = "tbl_MarketData") -> None:
    target_ws = None
    for ws in wb.worksheets:
        if table_name in ws.tables:
            target_ws = ws
            break

    if target_ws is None:
        target_ws = wb["Market_Data"] if "Market_Data" in wb.sheetnames else wb.create_sheet("Market_Data")

    for row in target_ws.iter_rows(
        min_row=1, max_row=max(target_ws.max_row, 1),
        min_col=1, max_col=max(target_ws.max_column, 1),
    ):
        for cell in row:
            cell.value = None

    if table_name in target_ws.tables:
        del target_ws.tables[table_name]

    for ci, h in enumerate(MARKET_HEADERS, 1):
        cell = target_ws.cell(row=1, column=ci, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")

    if market_df.empty:
        for ci in range(1, len(MARKET_HEADERS) + 1):
            target_ws.cell(row=2, column=ci, value="")
        last_row = 2
    else:
        for ri, (_, r) in enumerate(market_df.iterrows(), 2):
            for ci, h in enumerate(MARKET_HEADERS, 1):
                target_ws.cell(row=ri, column=ci, value=r.get(h))
        last_row = len(market_df) + 1

    ref = f"A1:{get_column_letter(len(MARKET_HEADERS))}{last_row}"
    table = Table(displayName=table_name, ref=ref)
    table.tableStyleInfo = TABLE_STYLE
    target_ws.add_table(table)

    for col in target_ws.columns:
        max_len = max(len(str(c.value or "")) for c in col)
        target_ws.column_dimensions[get_column_letter(col[0].column)].width = max(14, min(40, max_len + 4))

    target_ws.freeze_panes = "B2"
    target_ws.sheet_view.showGridLines = False
    target_ws.sheet_view.zoomScale = 85


def main() -> None:
    cfg = parse_args()
    if not cfg.workbook.exists():
        raise FileNotFoundError(f"Workbook not found: {cfg.workbook}")

    wb = load_workbook(cfg.workbook)
    trade_df, _, _ = read_table_dataframe(wb, "tbl_TradeLog")
    tickers, start_date, contract_hints = extract_tickers_start_and_hints(trade_df)
    if not tickers:
        write_market_table(wb, pd.DataFrame())
        wb.save(cfg.workbook)
        print(f"Updated tbl_MarketData for 0 tickers in {cfg.workbook}")
        return

    ib = IB()
    ib.connect(cfg.ib_host, cfg.ib_port, clientId=cfg.ib_client_id)

    all_data: list[pd.DataFrame] = []
    for ticker in tickers:
        series = pd.DataFrame()
        try:
            series = fetch_ib_series(ib, ticker, start_date, contract_hints.get(ticker))
        except Exception as exc:
            print(f"IB fetch failed for {ticker}: {exc}")

        if series.empty and should_use_yf(cfg, ticker):
            print(f"Using yfinance fallback for {ticker}")
            series = fetch_yf_series(ticker, start_date)

        if series.empty:
            print(f"No market data fetched for {ticker}")
            continue

        all_data.append(series)

    ib.disconnect()

    if not all_data:
        combined = pd.DataFrame()
    else:
        combined = pd.concat(all_data, ignore_index=True)
        combined = combined.sort_values(["Date", "Ticker"]).reset_index(drop=True)

    write_market_table(wb, combined)
    wb.save(cfg.workbook)
    print(f"Updated tbl_MarketData for {len(tickers)} tickers in {cfg.workbook}")


if __name__ == "__main__":
    main()
