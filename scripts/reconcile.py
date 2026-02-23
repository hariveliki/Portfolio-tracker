#!/usr/bin/env python3
"""Reconcile workbook holdings/cash against current IBKR account values."""

from __future__ import annotations

import argparse
import os
from typing import Any

import yaml
from dataclasses import dataclass
from pathlib import Path

from ib_insync import IB
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.worksheet.table import Table, TableStyleInfo


DEFAULT_CONFIG = "config/portfolio.yaml"


@dataclass
class Row:
    kind: str
    item: str
    workbook: float
    ibkr: float
    diff: float
    tolerance: float
    within_tolerance: bool
    notes: str


def load_yaml_config(config_path: Path) -> dict[str, Any]:
    if not config_path.exists():
        return {}
    with config_path.open("r", encoding="utf-8") as handle:
        return yaml.safe_load(handle) or {}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook")
    parser.add_argument("--config", default=DEFAULT_CONFIG)
    parser.add_argument("--ib-host")
    parser.add_argument("--ib-port", type=int)
    parser.add_argument("--ib-client-id", type=int)
    parser.add_argument("--cash-tag", default=os.getenv("IB_CASH_TAG", "TotalCashValue"))
    parser.add_argument("--qty-tol", type=float, default=1e-6)
    parser.add_argument("--cash-tol", type=float, default=0.01)
    args = parser.parse_args()

    raw_cfg = load_yaml_config(Path(args.config))
    ib_cfg = raw_cfg.get("ib", {})
    path_cfg = raw_cfg.get("paths", {})

    args.workbook = args.workbook or path_cfg.get("workbook_output", "output/portfolio_workbook.xlsx")
    args.ib_host = args.ib_host or os.getenv("IB_HOST", str(ib_cfg.get("host", "127.0.0.1")))
    args.ib_port = args.ib_port or int(os.getenv("IB_PORT", str(ib_cfg.get("port", 7497))))
    args.ib_client_id = args.ib_client_id or int(os.getenv("IB_CLIENT_ID", str(ib_cfg.get("client_id", 16))))
    return args


def table_rows(wb: Workbook, table_name: str) -> list[dict]:
    for ws in wb.worksheets:
        if table_name in ws.tables:
            table = ws.tables[table_name]
            min_col, min_row, max_col, max_row = range_boundaries(table.ref)
            values = list(
                ws.iter_rows(
                    min_row=min_row,
                    max_row=max_row,
                    min_col=min_col,
                    max_col=max_col,
                    values_only=True,
                )
            )
            if len(values) < 2:
                return []
            headers = values[0]
            output = []
            for raw in values[1:]:
                if not any(v is not None for v in raw):
                    continue
                output.append({str(headers[i]): raw[i] for i in range(len(headers))})
            return output
    raise ValueError(f"Table '{table_name}' not found")


def read_workbook_holdings(wb: Workbook) -> dict[str, float]:
    rows = table_rows(wb, "tbl_Positions")
    holdings: dict[str, float] = {}
    for row in rows:
        ticker = str(row.get("Ticker", "")).strip().upper()
        if not ticker:
            continue
        qty = row.get("Units", row.get("Quantity", row.get("Qty", row.get("Shares", 0)))) or 0
        holdings[ticker] = float(qty)
    return holdings


def read_workbook_cash(wb: Workbook) -> float:
    rows = table_rows(wb, "tbl_DailyCash")
    dated_rows: list[tuple[str, float]] = []
    for row in rows:
        date_val = row.get("Date")
        cash_val = row.get("Cash_Balance")
        if date_val is None or cash_val is None:
            continue
        dated_rows.append((str(date_val), float(cash_val)))

    if not dated_rows:
        raise ValueError("tbl_DailyCash has no populated Date/Cash_Balance rows")

    dated_rows.sort(key=lambda item: item[0])
    return dated_rows[-1][1]


def ib_positions_and_cash(args: argparse.Namespace) -> tuple[dict[str, float], float]:
    ib = IB()
    ib.connect(args.ib_host, args.ib_port, clientId=args.ib_client_id)

    positions: dict[str, float] = {}
    for p in ib.positions():
        ticker = getattr(p.contract, "symbol", "")
        if ticker:
            positions[ticker.upper()] = positions.get(ticker.upper(), 0.0) + float(p.position)

    cash_value = None
    for item in ib.accountSummary():
        if item.tag == args.cash_tag:
            cash_value = float(item.value)
            break

    ib.disconnect()

    if cash_value is None:
        raise ValueError(f"Cash tag '{args.cash_tag}' not found in account summary")
    return positions, cash_value


def write_reconciliation_sheet(wb: Workbook, rows: list[Row]) -> None:
    ws = wb["Reconciliation"] if "Reconciliation" in wb.sheetnames else wb.create_sheet("Reconciliation")

    for row in ws.iter_rows(min_row=1, max_row=max(ws.max_row, 1), min_col=1, max_col=max(ws.max_column, 1)):
        for cell in row:
            cell.value = None

    headers = ["Type", "Item", "Workbook", "IBKR", "Diff", "Tolerance", "WithinTolerance", "Notes"]
    ws.append(headers)
    for r in rows:
        ws.append([r.kind, r.item, r.workbook, r.ibkr, r.diff, r.tolerance, "Y" if r.within_tolerance else "N", r.notes])

    ref = f"A1:{get_column_letter(len(headers))}{max(len(rows)+1,2)}"
    if "tbl_Reconciliation" in ws.tables:
        ws.tables["tbl_Reconciliation"].ref = ref
    else:
        table = Table(displayName="tbl_Reconciliation", ref=ref)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
        ws.add_table(table)


def main() -> None:
    args = parse_args()
    workbook_path = Path(args.workbook)
    if not workbook_path.exists():
        raise FileNotFoundError(f"Workbook not found: {workbook_path}")

    wb = load_workbook(workbook_path)
    wb_holdings = read_workbook_holdings(wb)
    wb_cash = read_workbook_cash(wb)

    ib_holdings, ib_cash = ib_positions_and_cash(args)

    rows: list[Row] = []
    for ticker in sorted(set(wb_holdings) | set(ib_holdings)):
        w_qty = wb_holdings.get(ticker, 0.0)
        i_qty = ib_holdings.get(ticker, 0.0)
        diff = w_qty - i_qty
        rows.append(
            Row(
                kind="Position",
                item=ticker,
                workbook=w_qty,
                ibkr=i_qty,
                diff=diff,
                tolerance=args.qty_tol,
                within_tolerance=abs(diff) <= args.qty_tol,
                notes="" if abs(diff) <= args.qty_tol else "Quantity mismatch",
            )
        )

    cash_diff = wb_cash - ib_cash
    rows.append(
        Row(
            kind="Cash",
            item=args.cash_tag,
            workbook=wb_cash,
            ibkr=ib_cash,
            diff=cash_diff,
            tolerance=args.cash_tol,
            within_tolerance=abs(cash_diff) <= args.cash_tol,
            notes="" if abs(cash_diff) <= args.cash_tol else "Cash mismatch",
        )
    )

    write_reconciliation_sheet(wb, rows)
    wb.save(workbook_path)
    print(f"Reconciliation updated in {workbook_path}")


if __name__ == "__main__":
    main()
