from __future__ import annotations

from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.table import Table, TableStyleInfo

DEFAULT_CONFIG_PATH = Path("config/portfolio.yaml")
DEFAULT_OUTPUT_PATH = Path("output/portfolio_workbook.xlsx")

HEADER_FILL = PatternFill(start_color="1C2541", end_color="1C2541", fill_type="solid")
HEADER_FONT = Font(name="Segoe UI", color="FFFFFF", bold=True, size=10)
GREEN_FONT = Font(name="Segoe UI", color="007A33", size=10)
RED_FONT = Font(name="Segoe UI", color="B81D13", size=10)
GREEN_FILL = PatternFill(start_color="E2F0D9", end_color="E2F0D9", fill_type="solid")
RED_FILL = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
KPI_LABEL_FONT = Font(name="Segoe UI", bold=True, size=11, color="1C2541")
KPI_VALUE_FONT = Font(name="Segoe UI", size=11)
TITLE_FONT = Font(name="Segoe UI", bold=True, size=14, color="1C2541")

TABLE_STYLE = TableStyleInfo(
    name="TableStyleMedium2",
    showFirstColumn=False,
    showLastColumn=False,
    showRowStripes=True,
    showColumnStripes=False,
)

SHEETS = [
    "Dashboard",
    "Positions",
    "Equity_Curve",
    "Sector_Exposure",
    "Geographic_Exposure",
    "Trade_Log",
    "Market_Data",
    "Reference_Data",
    "Daily_Units",
    "Daily_Cash",
    "Settings",
]


def load_output_path(config_path: Path = DEFAULT_CONFIG_PATH) -> Path:
    if not config_path.exists():
        return DEFAULT_OUTPUT_PATH
    with config_path.open("r", encoding="utf-8") as handle:
        raw = yaml.safe_load(handle) or {}
    return Path((raw.get("paths") or {}).get("workbook_output", str(DEFAULT_OUTPUT_PATH)))


def _apply_sheet_defaults(ws) -> None:
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 85


def _autosize_columns(ws) -> None:
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = max(14, min(40, max_len + 4))


def _style_header_row(ws) -> None:
    for cell in ws[1]:
        if cell.value is None:
            continue
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")


def _make_table(ws, table_name: str, headers: list[str], seed_row: list | None = None) -> Table:
    ws.append(headers)
    ws.append(seed_row if seed_row else [""] * len(headers))
    ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"
    table = Table(displayName=table_name, ref=ref)
    table.tableStyleInfo = TABLE_STYLE
    ws.add_table(table)
    _style_header_row(ws)
    _autosize_columns(ws)
    ws.freeze_panes = "B2"
    _apply_sheet_defaults(ws)
    return table


def _apply_pos_neg_cf(ws, col_letters: list[str], max_row: int = 5000) -> None:
    for letter in col_letters:
        rng = f"{letter}2:{letter}{max_row}"
        ws.conditional_formatting.add(
            rng, CellIsRule(operator="greaterThan", formula=["0"], fill=GREEN_FILL, font=GREEN_FONT),
        )
        ws.conditional_formatting.add(
            rng, CellIsRule(operator="lessThan", formula=["0"], fill=RED_FILL, font=RED_FONT),
        )


def build_workbook(output_path: Path | None = None) -> Path:
    output_path = output_path or load_output_path()
    wb = Workbook()
    wb.remove(wb.active)

    for name in SHEETS:
        wb.create_sheet(name)

    ws = wb["Settings"]
    ws["A1"] = "Parameter"
    ws["B1"] = "Value"
    ws["A2"] = "Start Capital"
    ws["B2"] = 100000
    ws["A3"] = "Start Date"
    ws["B3"] = "2024-01-01"
    ws["A4"] = "Risk Free Rate"
    ws["B4"] = 0.02
    ws["A5"] = "Base Currency"
    ws["B5"] = "USD"
    _style_header_row(ws)
    _autosize_columns(ws)
    ws.freeze_panes = "B2"
    _apply_sheet_defaults(ws)

    wb.defined_names.add(DefinedName("StartCapital", attr_text="Settings!$B$2"))
    wb.defined_names.add(DefinedName("StartDate", attr_text="Settings!$B$3"))
    wb.defined_names.add(DefinedName("RiskFreeRate", attr_text="Settings!$B$4"))
    wb.defined_names.add(DefinedName("BaseCcy", attr_text="Settings!$B$5"))

    _make_table(
        wb["Trade_Log"],
        "tbl_TradeLog",
        [
            "Date", "IB_ExecId", "Ticker", "Action", "Shares", "Price", "Commission",
            "permId", "account", "exchange", "currency",
            "Gross_Value", "Net_Cost", "Signed_Shares", "Signed_Cash", "Cash_Balance",
        ],
        [
            "", "", "", "", 0, 0, 0, "", "", "", "",
            "=[@Shares]*[@Price]",
            '=IF(OR(UPPER([@Action])="BUY",UPPER([@Action])="BOT"),[@Gross_Value]+[@Commission],[@Gross_Value]-[@Commission])',
            '=IF(OR(UPPER([@Action])="BUY",UPPER([@Action])="BOT"),[@Shares],-[@Shares])',
            "=-[@Signed_Shares]*[@Price]-[@Commission]",
            "=StartCapital+SUM(INDEX(tbl_TradeLog[Signed_Cash],1):[@Signed_Cash])",
        ],
    )
    _apply_pos_neg_cf(wb["Trade_Log"], ["O"])

    _make_table(
        wb["Market_Data"],
        "tbl_MarketData",
        ["Date", "Ticker", "Open", "High", "Low", "Close", "Adj_Close", "Volume"],
    )

    _make_table(
        wb["Reference_Data"],
        "tbl_ReferenceData",
        ["Ticker", "Name", "Sector", "Industry", "Region", "Currency", "Beta"],
    )

    _make_table(
        wb["Daily_Units"],
        "tbl_DailyUnits",
        ["Date", "Ticker", "Units"],
    )

    _make_table(
        wb["Daily_Cash"],
        "tbl_DailyCash",
        ["Date", "Cash_Balance"],
    )

    _make_table(
        wb["Positions"],
        "tbl_Positions",
        [
            "Ticker", "Company", "Sector", "Region", "Units", "Avg_Cost", "Last_Price",
            "Day_Change_Pct", "Cost_Basis", "Market_Value", "Unrealized_PnL", "PnL_Pct", "Weight",
        ],
        [
            "", "", "", "", 0, 0, 0, 0,
            "=[@Units]*[@Avg_Cost]",
            "=[@Units]*[@Last_Price]",
            "=[@Market_Value]-[@Cost_Basis]",
            "=IFERROR([@Unrealized_PnL]/[@Cost_Basis],0)",
            "=IFERROR([@Market_Value]/SUM(tbl_Positions[Market_Value]),0)",
        ],
    )
    _apply_pos_neg_cf(wb["Positions"], ["H", "K", "L"])

    _make_table(
        wb["Equity_Curve"],
        "tbl_EquityCurve",
        [
            "Date", "Portfolio_Value", "Cash_Balance", "NAV", "Daily_Return",
            "Cumulative_Return", "Peak_NAV", "Drawdown",
        ],
        [
            "", 0, 0,
            "=[@Portfolio_Value]+[@Cash_Balance]",
            "=IFERROR([@NAV]/OFFSET([@NAV],-1,0)-1,0)",
            "=IFERROR([@NAV]/INDEX(tbl_EquityCurve[NAV],1)-1,0)",
            "=MAX(INDEX(tbl_EquityCurve[NAV],1):[@NAV])",
            "=IFERROR([@NAV]/[@Peak_NAV]-1,0)",
        ],
    )
    _apply_pos_neg_cf(wb["Equity_Curve"], ["E", "F", "H"])

    _make_table(
        wb["Sector_Exposure"],
        "tbl_SectorExposure",
        ["Sector", "Market_Value", "Weight"],
        [
            "",
            '=SUMIFS(tbl_Positions[Market_Value],tbl_Positions[Sector],[@Sector])',
            "=IFERROR([@Market_Value]/SUM(tbl_SectorExposure[Market_Value]),0)",
        ],
    )

    _make_table(
        wb["Geographic_Exposure"],
        "tbl_GeoExposure",
        ["Region", "Market_Value", "Weight"],
        [
            "",
            '=SUMIFS(tbl_Positions[Market_Value],tbl_Positions[Region],[@Region])',
            "=IFERROR([@Market_Value]/SUM(tbl_GeoExposure[Market_Value]),0)",
        ],
    )

    ws = wb["Dashboard"]
    _apply_sheet_defaults(ws)
    ws.freeze_panes = "A2"

    ws.merge_cells("B2:E2")
    ws["B2"] = "Portfolio Dashboard"
    ws["B2"].font = TITLE_FONT
    ws["B2"].alignment = Alignment(horizontal="left")

    kpis = [
        ("B4", "NAV", "C4", "=IFERROR(TAKE(tbl_EquityCurve[NAV],-1),0)", "#,##0.00"),
        ("B5", "Total Return", "C5", "=IFERROR(TAKE(tbl_EquityCurve[Cumulative_Return],-1),0)", "0.00%"),
        ("B6", "Daily Return", "C6", "=IFERROR(TAKE(tbl_EquityCurve[Daily_Return],-1),0)", "0.00%"),
        (
            "B7", "Sharpe Ratio", "C7",
            "=IFERROR((AVERAGE(tbl_EquityCurve[Daily_Return])*252-RiskFreeRate)"
            "/(STDEV(tbl_EquityCurve[Daily_Return])*SQRT(252)),0)",
            "0.00",
        ),
        ("B8", "Max Drawdown", "C8", "=IFERROR(MIN(tbl_EquityCurve[Drawdown]),0)", "0.00%"),
        ("B9", "VaR (95% Daily)", "C9", "=IFERROR(PERCENTILE(tbl_EquityCurve[Daily_Return],0.05),0)", "0.00%"),
        ("B10", "# Positions", "C10", '=COUNTIF(tbl_Positions[Units],">"&0)', "0"),
        ("B11", "Start Date", "C11", "=StartDate", "@"),
    ]
    for label_cell, label, value_cell, formula, fmt in kpis:
        ws[label_cell] = label
        ws[label_cell].font = KPI_LABEL_FONT
        ws[value_cell] = formula
        ws[value_cell].font = KPI_VALUE_FONT
        ws[value_cell].number_format = fmt

    ws["E4"] = "Correlation Matrix"
    ws["E4"].font = KPI_LABEL_FONT

    ws.column_dimensions["A"].width = 2
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 18

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path


if __name__ == "__main__":
    saved = build_workbook()
    print(f"Workbook saved to: {saved}")
