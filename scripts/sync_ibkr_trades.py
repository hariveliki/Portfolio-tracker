#!/usr/bin/env python3
"""Sync IBKR executions into the TradeLog workbook via Client Portal REST API."""

from __future__ import annotations

import json
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

import yaml
from openpyxl import load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.worksheet.table import Table

from ibkr_client import IbkrClient

RAW_COLUMNS = [
    "IB_ExecId",
    "Date",
    "Ticker",
    "Action",
    "Shares",
    "Price",
    "Commission",
    "permId",
    "account",
    "exchange",
    "currency",
]
FORMULA_COLUMNS = ["Gross_Value", "Net_Cost", "Signed_Shares", "Signed_Cash", "Cash_Balance"]


@dataclass
class SyncConfig:
    gateway_url: str
    account_id: str
    workbook_path: Path


def load_config(config_path: Path) -> SyncConfig:
    with config_path.open("r", encoding="utf-8") as handle:
        raw = yaml.safe_load(handle) or {}

    cpapi = raw.get("cpapi", {})
    paths = raw.get("paths", {})

    return SyncConfig(
        gateway_url=str(cpapi.get("gateway_url", "https://localhost:5001/v1/api")),
        account_id=str(cpapi.get("account_id", "")),
        workbook_path=Path(paths.get("workbook_output", "output/portfolio_workbook.xlsx")),
    )


def load_state(state_path: Path) -> dict[str, Any]:
    if not state_path.exists():
        return {}
    with state_path.open("r", encoding="utf-8") as handle:
        return json.load(handle)


def save_state(state_path: Path, state: dict[str, Any]) -> None:
    state_path.parent.mkdir(parents=True, exist_ok=True)
    with state_path.open("w", encoding="utf-8") as handle:
        json.dump(state, handle, indent=2, sort_keys=True)
        handle.write("\n")


def parse_trade_time(trade_time: str) -> str:
    try:
        dt = datetime.strptime(trade_time, "%Y%m%d-%H:%M:%S")
        return dt.date().isoformat()
    except ValueError:
        return datetime.now(tz=UTC).date().isoformat()


def find_trade_table(workbook_path: Path, table_name: str = "tbl_TradeLog") -> tuple[Any, Worksheet, Table]:
    wb = load_workbook(workbook_path)
    for ws in wb.worksheets:
        for name, table in ws.tables.items():
            if name == table_name:
                return wb, ws, table
    wb.close()
    raise RuntimeError(f"Could not find table {table_name!r} in {workbook_path}")


def read_table_headers(ws: Worksheet, table: Table) -> tuple[list[str], int, int, int, int]:
    start, end = table.ref.split(":")
    start_col = ws[start].column
    start_row = ws[start].row
    end_col = ws[end].column
    end_row = ws[end].row
    headers = [ws.cell(row=start_row, column=col).value for col in range(start_col, end_col + 1)]
    return [str(h) if h is not None else "" for h in headers], start_col, start_row, end_col, end_row


def existing_exec_ids(ws: Worksheet, headers: list[str], start_col: int, start_row: int, end_row: int) -> set[str]:
    if "IB_ExecId" not in headers:
        raise RuntimeError("tbl_TradeLog must include IB_ExecId column")
    idx = headers.index("IB_ExecId")
    col = start_col + idx
    values: set[str] = set()
    for row in range(start_row + 1, end_row + 1):
        cell_value = ws.cell(row=row, column=col).value
        if cell_value is not None and str(cell_value).strip():
            values.add(str(cell_value).strip())
    return values


def map_trade(trade: dict[str, Any]) -> dict[str, Any]:
    return {
        "IB_ExecId": trade.get("execution_id", ""),
        "Date": parse_trade_time(trade.get("trade_time", "")),
        "Ticker": trade.get("symbol", ""),
        "Action": trade.get("side", ""),
        "Shares": trade.get("size", 0),
        "Price": float(trade.get("price", 0)),
        "Commission": float(trade.get("commission", 0)),
        "permId": None,
        "account": trade.get("account", ""),
        "exchange": trade.get("exchange", ""),
        "currency": None,
    }


def copy_formula_columns(
    ws: Worksheet,
    headers: list[str],
    start_col: int,
    template_row: int,
    first_new_row: int,
    last_new_row: int,
) -> None:
    for col_name in FORMULA_COLUMNS:
        if col_name not in headers:
            continue
        col = start_col + headers.index(col_name)
        template_cell = ws.cell(row=template_row, column=col)
        formula = template_cell.value
        if not isinstance(formula, str) or not formula.startswith("="):
            continue
        template_coordinate = template_cell.coordinate
        for row in range(first_new_row, last_new_row + 1):
            new_cell = ws.cell(row=row, column=col)
            new_cell.value = Translator(formula, origin=template_coordinate).translate_formula(new_cell.coordinate)


def append_rows(workbook_path: Path, rows: list[dict[str, Any]]) -> int:
    if not rows:
        return 0

    wb, ws, table = find_trade_table(workbook_path)
    headers, start_col, start_row, end_col, end_row = read_table_headers(ws, table)

    first_new_row = end_row + 1
    for row_idx, row_data in enumerate(rows, start=first_new_row):
        for col_offset, header in enumerate(headers):
            col = start_col + col_offset
            if header in RAW_COLUMNS:
                ws.cell(row=row_idx, column=col).value = row_data.get(header)

    copy_formula_columns(
        ws,
        headers,
        start_col,
        template_row=end_row,
        first_new_row=first_new_row,
        last_new_row=first_new_row + len(rows) - 1,
    )

    table.ref = f"{ws.cell(row=start_row, column=start_col).coordinate}:{ws.cell(row=first_new_row + len(rows) - 1, column=end_col).coordinate}"
    wb.save(workbook_path)
    wb.close()
    return len(rows)


def main() -> None:
    root = Path(__file__).resolve().parents[1]
    config = load_config(root / "config" / "portfolio.yaml")
    state_path = root / "state" / "state.json"
    state = load_state(state_path)

    with IbkrClient(base_url=config.gateway_url, account_id=config.account_id) as client:
        client.check_auth()

        trades = client.get_trades(days=7)
        mapped = [map_trade(t) for t in trades]

        wb, ws, table = find_trade_table(config.workbook_path)
        headers, start_col, start_row, _, end_row = read_table_headers(ws, table)
        known_ids = existing_exec_ids(ws, headers, start_col, start_row, end_row)
        wb.close()

        new_rows: list[dict[str, Any]] = []
        max_epoch: int = 0
        for trade, row in zip(trades, mapped):
            exec_id = str(row["IB_ExecId"])
            if exec_id in known_ids:
                continue
            new_rows.append(row)
            epoch = trade.get("trade_time_r", 0)
            if isinstance(epoch, int) and epoch > max_epoch:
                max_epoch = epoch

        inserted = append_rows(config.workbook_path, new_rows)

        if inserted > 0:
            if max_epoch:
                ts = datetime.fromtimestamp(max_epoch / 1000, tz=UTC)
            else:
                ts = datetime.now(tz=UTC)
            state["last_sync_utc"] = ts.isoformat().replace("+00:00", "Z")
            save_state(state_path, state)


if __name__ == "__main__":
    main()
