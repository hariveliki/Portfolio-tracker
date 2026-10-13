#!/usr/bin/env python3
"""Rebuild all analytics tables (Daily_Units, Daily_Cash, Positions, Equity_Curve,
Sector/Geographic/Currency Exposure, Contribution, Rebalancing, Dashboard) from the data-lake sheets."""

from __future__ import annotations

import argparse
import math
from datetime import date
from pathlib import Path
from typing import Any

import pandas as pd
import yaml
import yfinance as yf
from openpyxl import load_workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference as ChartRef
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import DataPoint
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.worksheet.table import Table, TableStyleInfo

DEFAULT_CONFIG = "config/portfolio.yaml"

HEADER_FILL = PatternFill(start_color="1C2541", end_color="1C2541", fill_type="solid")
HEADER_FONT = Font(name="Segoe UI", color="FFFFFF", bold=True, size=10)
GREEN_FONT = Font(name="Segoe UI", color="007A33", size=10)
RED_FONT = Font(name="Segoe UI", color="B81D13", size=10)
GREEN_FILL = PatternFill(start_color="E2F0D9", end_color="E2F0D9", fill_type="solid")
RED_FILL = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
INPUT_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
KPI_LABEL_FONT = Font(name="Segoe UI", bold=True, size=11, color="1C2541")
KPI_VALUE_FONT = Font(name="Segoe UI", size=11)
TITLE_FONT = Font(name="Segoe UI", bold=True, size=14, color="1C2541")

SIGNED_PCT = "+0.00%;-0.00%;0.00%"
NAV_CHART_HEIGHT_CM = 12
ROW_HEIGHT_CM = 0.53  # default Excel row height (15 pt)

TABLE_STYLE = TableStyleInfo(
    name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
    showRowStripes=True, showColumnStripes=False,
)


def _rows_for(height_cm: float) -> int:
    return math.ceil(height_cm / ROW_HEIGHT_CM)


def load_yaml_config(config_path: Path) -> dict[str, Any]:
    if not config_path.exists():
        return {}
    with config_path.open("r", encoding="utf-8") as handle:
        return yaml.safe_load(handle) or {}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook")
    parser.add_argument("--config", default=DEFAULT_CONFIG)
    args = parser.parse_args()
    raw_cfg = load_yaml_config(Path(args.config))
    path_cfg = raw_cfg.get("paths", {})
    args.workbook = args.workbook or path_cfg.get("workbook_output", "output/portfolio_workbook.xlsx")
    return args


def _read_table_df(wb, table_name: str) -> pd.DataFrame:
    for ws in wb.worksheets:
        if table_name not in ws.tables:
            continue
        table = ws.tables[table_name]
        min_col, min_row, max_col, max_row = range_boundaries(table.ref)
        values = list(ws.iter_rows(
            min_row=min_row, max_row=max_row,
            min_col=min_col, max_col=max_col,
            values_only=True,
        ))
        if len(values) < 2:
            return pd.DataFrame()
        headers = [str(h) for h in values[0]]
        rows = [
            r for r in values[1:]
            if any(
                v is not None and str(v).strip() and not str(v).startswith("=")
                for v in r
            )
        ]
        return pd.DataFrame(rows, columns=headers)
    return pd.DataFrame()


def _read_setting(wb, param_name: str, default: Any = None) -> Any:
    if "Settings" not in wb.sheetnames:
        return default
    ws = wb["Settings"]
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=2, values_only=True):
        if row[0] and str(row[0]).strip() == param_name:
            return row[1]
    return default


def _clear_sheet(ws, max_col_override: int | None = None) -> None:
    for merge in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(merge))
    max_col = max_col_override or max(ws.max_column, 1)
    for row in ws.iter_rows(min_row=1, max_row=max(ws.max_row, 1), min_col=1, max_col=max_col):
        for cell in row:
            cell.value = None
    for tname in list(ws.tables.keys()):
        del ws.tables[tname]
    ws.conditional_formatting._cf_rules = {}


def _write_table(ws, table_name: str, headers: list[str], rows: list[list],
                 number_formats: dict[int, str] | None = None) -> None:
    for ci, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=ci, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")

    if not rows:
        for ci in range(1, len(headers) + 1):
            ws.cell(row=2, column=ci, value="")
        last_row = 2
    else:
        for ri, row_data in enumerate(rows, 2):
            for ci, val in enumerate(row_data, 1):
                cell = ws.cell(row=ri, column=ci, value=val)
                if number_formats and ci in number_formats:
                    cell.number_format = number_formats[ci]
        last_row = len(rows) + 1

    ref = f"A1:{get_column_letter(len(headers))}{last_row}"
    table = Table(displayName=table_name, ref=ref)
    table.tableStyleInfo = TABLE_STYLE
    ws.add_table(table)

    for col in ws.columns:
        max_len = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = max(14, min(40, max_len + 4))

    ws.freeze_panes = "B2"
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 85


def _apply_pos_neg_cf(ws, col_letters: list[str], max_row: int) -> None:
    for letter in col_letters:
        rng = f"{letter}2:{letter}{max_row}"
        ws.conditional_formatting.add(
            rng, CellIsRule(operator="greaterThan", formula=["0"], fill=GREEN_FILL, font=GREEN_FONT),
        )
        ws.conditional_formatting.add(
            rng, CellIsRule(operator="lessThan", formula=["0"], fill=RED_FILL, font=RED_FONT),
        )


def _prepare_trades(raw: pd.DataFrame) -> pd.DataFrame:
    if raw.empty:
        return pd.DataFrame(columns=["Date", "Ticker", "Shares", "Price", "Commission", "Action",
                                     "Signed_Shares", "Signed_Cash"])
    df = raw.copy()
    df["Shares"] = pd.to_numeric(df["Shares"], errors="coerce").fillna(0)
    df["Price"] = pd.to_numeric(df["Price"], errors="coerce").fillna(0)
    df["Commission"] = pd.to_numeric(df["Commission"], errors="coerce").fillna(0)
    df["Date"] = pd.to_datetime(df["Date"], errors="coerce")
    df = df.dropna(subset=["Date"])
    df["Date"] = df["Date"].dt.date
    df["Ticker"] = df["Ticker"].astype(str).str.strip().str.upper()
    df = df[df["Ticker"].ne("")]
    df["Signed_Shares"] = df.apply(
        lambda r: r["Shares"] if str(r["Action"]).upper() in ("BUY", "BOT") else -r["Shares"], axis=1,
    )
    df["Signed_Cash"] = -df["Signed_Shares"] * df["Price"] - df["Commission"]
    return df


def _prepare_market(raw: pd.DataFrame) -> pd.DataFrame:
    if raw.empty:
        return pd.DataFrame(columns=["Date", "Ticker", "Close"])
    df = raw.copy()
    df["Date"] = pd.to_datetime(df["Date"], errors="coerce")
    df = df.dropna(subset=["Date"])
    df["Date"] = df["Date"].dt.date
    df["Ticker"] = df["Ticker"].astype(str).str.strip().str.upper()
    df["Close"] = pd.to_numeric(df["Close"], errors="coerce")
    return df


def _prepare_reference(raw: pd.DataFrame) -> pd.DataFrame:
    if raw.empty:
        return pd.DataFrame(columns=["Ticker", "Name", "Sector", "Industry", "Region", "Currency", "Beta"])
    df = raw.copy()
    df["Ticker"] = df["Ticker"].astype(str).str.strip().str.upper()
    return df


def _fx_rate(currency: str, base: str, dates: list[date]) -> pd.Series:
    """Value of one `currency` unit in `base` on each date (closed-market days carry over)."""
    hist = yf.download(f"{currency}{base}=X", start=min(dates), end=max(dates) + pd.Timedelta(days=5),
                       progress=False, auto_adjust=True)["Close"].squeeze()
    hist.index = hist.index.date
    return hist.reindex(sorted(set(hist.index) | set(dates))).ffill().bfill().reindex(dates)


def _ticker_currencies(trades: pd.DataFrame) -> pd.Series:
    """Listing currency per ticker: the last non-blank Trade_Log currency (empty if there is none)."""
    if trades.empty or "currency" not in trades.columns:
        return pd.Series(dtype=str)
    ccy = trades["currency"].fillna("").astype(str).str.upper()
    return trades.assign(_c=ccy).query("_c != ''").groupby("Ticker")["_c"].last()


def convert_to_base(trades: pd.DataFrame, market: pd.DataFrame, base: str) -> None:
    """Convert trade prices and closes from each ticker's currency to `base`, in place."""
    if trades.empty or "currency" not in trades.columns:
        return
    ticker_ccy = _ticker_currencies(trades)
    for currency in sorted(set(ticker_ccy) - {base}):
        tickers = ticker_ccy.index[ticker_ccy == currency]
        t_mask = trades["Ticker"].isin(tickers)
        m_mask = market["Ticker"].isin(tickers)
        dates = sorted(set(trades.loc[t_mask, "Date"]) | set(market.loc[m_mask, "Date"]))
        try:
            rate = _fx_rate(currency, base, dates)
        except Exception as exc:
            print(f"Warning: no {currency}->{base} rates, {', '.join(tickers)} stay in {currency}: {exc}")
            continue
        trades.loc[t_mask, "Price"] *= trades.loc[t_mask, "Date"].map(rate)
        market.loc[m_mask, "Close"] *= market.loc[m_mask, "Date"].map(rate)
    trades["Signed_Cash"] = -trades["Signed_Shares"] * trades["Price"] - trades["Commission"]


def compute_daily_units(trades: pd.DataFrame, trading_dates: list[date]) -> pd.DataFrame:
    if trades.empty or not trading_dates:
        return pd.DataFrame(columns=["Date", "Ticker", "Units"])

    trade_events = trades.groupby(["Date", "Ticker"])["Signed_Shares"].sum().reset_index()
    tickers = sorted(trade_events["Ticker"].unique())
    first_td = trading_dates[0]

    results: list[dict] = []
    for ticker in tickers:
        t_events = trade_events[trade_events["Ticker"] == ticker].set_index("Date")["Signed_Shares"]
        cumulative = float(t_events[t_events.index < first_td].sum())
        for d in trading_dates:
            if d in t_events.index:
                cumulative += t_events[d]
            if cumulative != 0:
                results.append({"Date": d, "Ticker": ticker, "Units": cumulative})

    return pd.DataFrame(results)


def compute_daily_cash(trades: pd.DataFrame, trading_dates: list[date], start_capital: float) -> pd.DataFrame:
    if not trading_dates:
        return pd.DataFrame(columns=["Date", "Cash_Balance"])

    cash_events = trades.groupby("Date")["Signed_Cash"].sum() if not trades.empty else pd.Series(dtype=float)

    balance = start_capital
    if not cash_events.empty:
        balance += float(cash_events[cash_events.index < trading_dates[0]].sum())

    rows: list[dict] = []
    for d in trading_dates:
        if d in cash_events.index:
            balance += cash_events[d]
        rows.append({"Date": d, "Cash_Balance": balance})
    return pd.DataFrame(rows)


def compute_positions(trades: pd.DataFrame, market: pd.DataFrame,
                      reference: pd.DataFrame) -> pd.DataFrame:
    if trades.empty:
        return pd.DataFrame()

    current_units = trades.groupby("Ticker")["Signed_Shares"].sum()
    current_units = current_units[current_units.abs() > 1e-9]
    if current_units.empty:
        return pd.DataFrame()

    avg_cost: dict[str, float] = {}
    for ticker in current_units.index:
        t = trades[trades["Ticker"] == ticker]
        total_cost = (t["Signed_Shares"] * t["Price"] + t["Commission"]).sum()
        total_shares = t["Signed_Shares"].sum()
        avg_cost[ticker] = abs(total_cost / total_shares) if total_shares != 0 else 0

    ref_map = {}
    if not reference.empty:
        for _, r in reference.iterrows():
            ref_map[r["Ticker"]] = r

    latest_prices: dict[str, float] = {}
    day_changes: dict[str, float] = {}
    if not market.empty:
        for ticker in current_units.index:
            t_market = market[market["Ticker"] == ticker].sort_values("Date")
            if t_market.empty:
                continue
            latest_prices[ticker] = float(t_market["Close"].iloc[-1])
            if len(t_market) >= 2:
                prev = float(t_market["Close"].iloc[-2])
                curr = float(t_market["Close"].iloc[-1])
                day_changes[ticker] = (curr / prev - 1) if prev != 0 else 0
            else:
                day_changes[ticker] = 0

    rows = []
    for ticker in sorted(current_units.index):
        units = current_units[ticker]
        ref = ref_map.get(ticker, {})
        last_price = latest_prices.get(ticker, 0)
        rows.append({
            "Ticker": ticker,
            "Company": (ref.get("Name") if isinstance(ref, (dict, pd.Series)) else None) or ticker,
            "Sector": ref.get("Sector", "") if isinstance(ref, (dict, pd.Series)) else "",
            "Region": ref.get("Region", "") if isinstance(ref, (dict, pd.Series)) else "",
            "Units": units,
            "Avg_Cost": avg_cost.get(ticker, 0),
            "Last_Price": last_price,
            "Day_Change_Pct": day_changes.get(ticker, 0),
        })
    return pd.DataFrame(rows)


def _filled_closes(market: pd.DataFrame, trading_dates: list[date]) -> pd.DataFrame:
    """Closes per trading date and ticker; exchanges have different holidays, so the last close carries over."""
    wide = market.pivot_table(index="Date", columns="Ticker", values="Close") if not market.empty else pd.DataFrame()
    return wide.reindex(trading_dates).ffill().bfill()


def fetch_benchmark(ticker: str, base: str, dates: list[date]) -> pd.Series:
    """Benchmark closes on `dates` (holidays carry over); all NaN, with a warning, if unavailable."""
    try:
        hist = yf.download(ticker, start=min(dates), end=max(dates) + pd.Timedelta(days=5),
                           progress=False, auto_adjust=True)["Close"].squeeze()
        if hist.empty:
            raise ValueError("download returned nothing")
        hist.index = hist.index.date
        closes = hist.reindex(sorted(set(hist.index) | set(dates))).ffill().bfill().reindex(dates)
    except Exception as exc:
        print(f"Warning: no benchmark data for {ticker}, benchmark cells stay empty: {exc}")
        return pd.Series(float("nan"), index=dates)
    try:
        currency = str(yf.Ticker(ticker).fast_info["currency"]).upper()
        if currency != base:
            print(f"Warning: benchmark {ticker} is quoted in {currency}, not {base}; no conversion applied")
    except Exception:
        pass  # currency unknown: nothing to compare
    return closes


def _has_benchmark(eq_curve: pd.DataFrame) -> bool:
    return "Benchmark_Close" in eq_curve and eq_curve["Benchmark_Close"].notna().any()


def compute_equity_curve(daily_units: pd.DataFrame, daily_cash: pd.DataFrame,
                         market: pd.DataFrame) -> pd.DataFrame:
    if daily_cash.empty:
        return pd.DataFrame(columns=["Date", "Portfolio_Value", "Cash_Balance"])

    trading_dates = sorted(daily_cash["Date"].unique())

    market_wide = _filled_closes(market, trading_dates)
    units_wide = daily_units.pivot_table(index="Date", columns="Ticker", values="Units", fill_value=0) if not daily_units.empty else pd.DataFrame()

    cash_lookup = dict(zip(daily_cash["Date"], daily_cash["Cash_Balance"]))

    rows: list[dict] = []
    for d in trading_dates:
        pv = 0.0
        if d in units_wide.index and d in market_wide.index:
            u_row = units_wide.loc[d]
            p_row = market_wide.loc[d]
            common = u_row.index.intersection(p_row.index)
            pv = float((u_row[common] * p_row[common]).sum())

        rows.append({
            "Date": d,
            "Portfolio_Value": round(pv, 2),
            "Cash_Balance": round(cash_lookup.get(d, 0), 2),
        })

    return pd.DataFrame(rows)


def _write_daily_units(wb, daily_units: pd.DataFrame) -> None:
    ws = wb["Daily_Units"]
    _clear_sheet(ws)
    rows = [[r["Date"], r["Ticker"], r["Units"]] for _, r in daily_units.iterrows()] if not daily_units.empty else []
    _write_table(ws, "tbl_DailyUnits", ["Date", "Ticker", "Units"], rows)


def _write_daily_cash(wb, daily_cash: pd.DataFrame) -> None:
    ws = wb["Daily_Cash"]
    _clear_sheet(ws)
    rows = [[r["Date"], round(r["Cash_Balance"], 2)] for _, r in daily_cash.iterrows()] if not daily_cash.empty else []
    _write_table(ws, "tbl_DailyCash", ["Date", "Cash_Balance"], rows,
                 number_formats={2: "#,##0.00"})


def _write_positions(wb, positions: pd.DataFrame) -> None:
    ws = wb["Positions"]
    _clear_sheet(ws)

    headers = [
        "Ticker", "Company", "Sector", "Region", "Units", "Avg_Cost", "Last_Price",
        "Day_Change_Pct", "Cost_Basis", "Market_Value", "Unrealized_PnL", "PnL_Pct", "Weight",
    ]
    rows = []
    for row_idx, (_, p) in enumerate(positions.iterrows(), start=2):
        rows.append([
            p["Ticker"], p["Company"], p["Sector"], p["Region"],
            p["Units"], round(p["Avg_Cost"], 4), round(p["Last_Price"], 4),
            p.get("Day_Change_Pct", 0),
            f"=E{row_idx}*F{row_idx}",
            f"=E{row_idx}*G{row_idx}",
            f"=J{row_idx}-I{row_idx}",
            f"=IFERROR(K{row_idx}/I{row_idx},0)",
            f"=IFERROR(J{row_idx}/SUM($J:$J),0)",
        ])

    fmt = {5: "#,##0.00", 6: "#,##0.00", 7: "#,##0.00", 8: "0.00%",
           9: "#,##0.00", 10: "#,##0.00", 11: "#,##0.00", 12: "0.00%", 13: "0.00%"}
    _write_table(ws, "tbl_Positions", headers, rows, number_formats=fmt)

    if rows:
        _apply_pos_neg_cf(ws, ["H", "K", "L"], len(rows) + 1)


def _write_equity_curve(wb, eq_curve: pd.DataFrame) -> None:
    ws = wb["Equity_Curve"]
    _clear_sheet(ws)

    headers = [
        "Date", "Portfolio_Value", "Cash_Balance", "NAV", "Daily_Return",
        "Cumulative_Return", "Peak_NAV", "Drawdown",
        "Benchmark_Close", "Benchmark_Return", "Benchmark_Cumulative",
    ]
    has_bench = _has_benchmark(eq_curve)
    rows = []
    for row_idx, (_, r) in enumerate(eq_curve.iterrows(), start=2):
        # Without benchmark data the three benchmark cells stay empty.
        bench = [r["Benchmark_Close"], f"=IFERROR(I{row_idx}/I{row_idx-1}-1,0)",
                 f"=IFERROR(I{row_idx}/$I$2-1,0)"] if has_bench else [None] * 3
        rows.append([
            r["Date"],
            r["Portfolio_Value"],
            r["Cash_Balance"],
            f"=B{row_idx}+C{row_idx}",
            f"=IFERROR(D{row_idx}/D{row_idx-1}-1,0)",
            f"=IFERROR(D{row_idx}/$D$2-1,0)",
            f"=MAX($D$2:D{row_idx})",
            f"=IFERROR(D{row_idx}/G{row_idx}-1,0)",
            *bench,
        ])

    fmt = {2: "#,##0.00", 3: "#,##0.00", 4: "#,##0.00",
           5: "0.00%", 6: "0.00%", 7: "#,##0.00", 8: "0.00%",
           9: "#,##0.00", 10: "0.00%", 11: "0.00%"}
    _write_table(ws, "tbl_EquityCurve", headers, rows, number_formats=fmt)

    if rows:
        _apply_pos_neg_cf(ws, ["E", "F", "H", "J", "K"], len(rows) + 1)


def compute_contribution(trades: pd.DataFrame, market: pd.DataFrame, daily_units: pd.DataFrame,
                         eq_curve: pd.DataFrame, reference: pd.DataFrame) -> pd.DataFrame:
    """Per-ticker PnL = End_Value - Start_Value + Net_Trades, so the PnLs add up to the NAV change."""
    columns = ["Ticker", "Company", "Start_Value", "End_Value", "Net_Trades", "PnL", "Contribution"]
    if eq_curve.empty or daily_units.empty:
        return pd.DataFrame(columns=columns)

    trading_dates = sorted(eq_curve["Date"])
    first, last = trading_dates[0], trading_dates[-1]
    nav_first = eq_curve["Portfolio_Value"].iloc[0] + eq_curve["Cash_Balance"].iloc[0]
    # A ticker without market data has no close: like the equity curve, value it at 0.
    close = _filled_closes(market, trading_dates).reindex(columns=sorted(daily_units["Ticker"].unique())).fillna(0)
    units = daily_units.pivot_table(index="Date", columns="Ticker", values="Units", fill_value=0)
    names = dict(zip(reference["Ticker"], reference["Name"])) if not reference.empty else {}
    # Like Daily_Cash, only trades on trading days after the first one move the NAV.
    later = trades[(trades["Date"] > first) & (trades["Date"] <= last) & trades["Date"].isin(trading_dates)]
    net_trades = later.groupby("Ticker")["Signed_Cash"].sum()

    rows = []
    for ticker in units.columns:
        start = float(units.loc[first, ticker]) * float(close.loc[first, ticker]) if first in units.index else 0.0
        end = float(units.loc[last, ticker]) * float(close.loc[last, ticker]) if last in units.index else 0.0
        net = float(net_trades.get(ticker, 0.0))
        pnl = end - start + net
        rows.append([ticker, names.get(ticker) or ticker, start, end, net, pnl, pnl / nav_first])
    return pd.DataFrame(rows, columns=columns).sort_values("Contribution", ascending=False)


def _write_contribution(wb, contribution: pd.DataFrame) -> None:
    ws = wb["Contribution"]
    _clear_sheet(ws)

    rows = []
    for row_idx, (_, c) in enumerate(contribution.iterrows(), start=2):
        rows.append([
            c["Ticker"], c["Company"], c["Start_Value"], c["End_Value"], c["Net_Trades"],
            f"=D{row_idx}-C{row_idx}+E{row_idx}",
            f"=IFERROR(F{row_idx}/Equity_Curve!$D$2,0)",
        ])

    fmt = {3: "#,##0.00", 4: "#,##0.00", 5: "#,##0.00", 6: "#,##0.00", 7: "0.00%"}
    _write_table(ws, "tbl_Contribution", ["Ticker", "Company", "Start_Value", "End_Value", "Net_Trades",
                                         "PnL", "Contribution"], rows, number_formats=fmt)
    if not rows:
        return
    _apply_pos_neg_cf(ws, ["F", "G"], len(rows) + 1)

    # The total sits right below the table, outside it, so sorting and the chart ignore it.
    total_row = len(rows) + 2
    ws.cell(row=total_row, column=1, value="Total").font = KPI_LABEL_FONT
    for col, letter in ((3, "C"), (4, "D"), (5, "E"), (6, "F"), (7, "G")):
        cell = ws.cell(row=total_row, column=col, value=f"=SUM({letter}2:{letter}{total_row - 1})")
        cell.font = KPI_LABEL_FONT
        cell.number_format = fmt[col]


def _write_exposure(wb, sheet: str, table_name: str, dimension: str, positions: pd.DataFrame,
                    exposure: pd.DataFrame) -> None:
    """Look-through exposure: sum of position value x fund weight per bucket."""
    ws = wb[sheet]
    _clear_sheet(ws)

    rows = []
    if not positions.empty and not exposure.empty:
        held = exposure[(exposure["Dimension"] == dimension) & exposure["Ticker"].isin(positions["Ticker"])]
        last = len(positions) + 1
        for row_idx, bucket in enumerate(sorted(held["Bucket"].unique()), start=2):
            rows.append([
                bucket,
                f"=SUMPRODUCT(Positions!$J$2:$J${last},SUMIFS(Reference_Exposure!$D:$D,"
                f"Reference_Exposure!$A:$A,Positions!$A$2:$A${last},"
                f'Reference_Exposure!$B:$B,"{dimension}",Reference_Exposure!$C:$C,A{row_idx}))',
                f"=IFERROR(B{row_idx}/SUM($B:$B),0)",
            ])

    _write_table(ws, table_name, [dimension, "Market_Value", "Weight"], rows,
                 number_formats={2: "#,##0.00", 3: "0.00%"})


def compute_currency_exposure(positions: pd.DataFrame, trades: pd.DataFrame) -> pd.Series:
    """Position value (units x last price, already in base currency) per listing currency."""
    if positions.empty:
        return pd.Series(dtype=float)
    value = positions["Units"] * positions["Last_Price"].round(4)  # Positions stores 4 decimals
    currency = positions["Ticker"].map(_ticker_currencies(trades)).fillna("Unknown")
    return value.groupby(currency).sum().sort_values(ascending=False)


def _write_currency_exposure(wb, by_currency: pd.Series) -> None:
    ws = wb["Currency_Exposure"]
    _clear_sheet(ws)

    rows = [[ccy, float(value), f"=IFERROR(B{row_idx}/SUM($B:$B),0)"]
            for row_idx, (ccy, value) in enumerate(by_currency.items(), start=2)]
    _write_table(ws, "tbl_CurrencyExposure", ["Currency", "Market_Value", "Weight"], rows,
                 number_formats={2: "#,##0.00", 3: "0.00%"})
    ws["E1"] = "Listing currency of each position, valued in base currency (not look-through exposure)"


def _read_targets(wb) -> dict[str, float]:
    """Target weights the user typed into Rebalancing, by ticker; blank or non-numeric cells are skipped."""
    df = _read_table_df(wb, "tbl_Rebalancing")
    if df.empty:
        return {}
    df["Target_Weight"] = pd.to_numeric(df["Target_Weight"], errors="coerce")
    df = df[df["Ticker"].str.strip().astype(bool)].dropna(subset=["Target_Weight"])
    return dict(zip(df["Ticker"], df["Target_Weight"].astype(float)))


def _write_rebalancing(wb, positions: pd.DataFrame, targets: dict[str, float]) -> int:
    """Targets are user input and `_clear_sheet` wipes them, so they are passed in and written back.
    Returns the number of table rows."""
    ws = wb["Rebalancing"]
    _clear_sheet(ws)

    held = positions["Ticker"].tolist() if not positions.empty else []
    tickers = held + sorted(t for t in targets if t not in held)  # saved target, no longer held
    value = "SUM(Positions!$J:$J)"
    rows = []
    for row_idx, ticker in enumerate(tickers, start=2):
        rows.append([
            ticker,
            f"=IFERROR(INDEX(Positions!$M:$M,MATCH(A{row_idx},Positions!$A:$A,0)),0)",
            targets.get(ticker),
            f'=IF(C{row_idx}="","",B{row_idx}-C{row_idx})',
            f'=IF(C{row_idx}="","",(C{row_idx}-B{row_idx})*{value})',
        ])

    _write_table(ws, "tbl_Rebalancing", ["Ticker", "Current_Weight", "Target_Weight", "Drift", "Trade_Value"],
                 rows, number_formats={2: "0.0%", 3: "0.0%", 4: "+0.0%;-0.0%;0.0%", 5: "#,##0.00"})
    ws["G1"] = "Current_Weight = share of invested securities, excluding cash"
    ws["G2"] = "Fill Target_Weight (e.g. 25%); targets should sum to 100%; positive Trade_Value = buy"
    if not rows:
        return 0

    last = len(rows) + 1
    for row_idx in range(2, last + 1):
        ws.cell(row=row_idx, column=3).fill = INPUT_FILL
    # Drift is "" without a target, and Excel ranks text above 0, so CellIsRule would colour it green.
    rng = f"D2:D{last}"
    ws.conditional_formatting.add(
        rng, FormulaRule(formula=["AND(ISNUMBER(D2),D2>0)"], fill=GREEN_FILL, font=GREEN_FONT))
    ws.conditional_formatting.add(
        rng, FormulaRule(formula=["AND(ISNUMBER(D2),D2<0)"], fill=RED_FILL, font=RED_FONT))

    # Total below the table, outside it, like the Contribution total.
    ws.cell(row=last + 1, column=1, value="Targets total").font = KPI_LABEL_FONT
    total = ws.cell(row=last + 1, column=3, value=f"=SUM(C2:C{last})")
    total.font = KPI_LABEL_FONT
    total.number_format = "0.0%"
    return len(rows)


def compute_total_commissions(trades: pd.DataFrame, eq_curve: pd.DataFrame) -> float:
    """Commissions that move the NAV: same trades as Net_Trades in compute_contribution."""
    if eq_curve.empty or trades.empty:
        return 0.0
    trading_dates = sorted(eq_curve["Date"])
    first, last = trading_dates[0], trading_dates[-1]
    later = trades[(trades["Date"] > first) & (trades["Date"] <= last) & trades["Date"].isin(trading_dates)]
    return float(later["Commission"].sum())


def _write_dashboard(wb, eq_curve: pd.DataFrame, positions: pd.DataFrame,
                     market: pd.DataFrame, contribution: pd.DataFrame, commissions: float,
                     bench_ticker: str, drift_rows: int) -> None:
    ws = wb["Dashboard"]
    _clear_sheet(ws, max_col_override=30)
    ws._charts.clear()

    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 85
    ws.freeze_panes = "A3"
    ws.column_dimensions["A"].width = 2
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 4

    ws.merge_cells("B2:E2")
    ws["B2"] = "Portfolio Dashboard"
    ws["B2"].font = TITLE_FONT
    ws["B2"].alignment = Alignment(horizontal="left")

    # Risk KPIs use daily returns from row 3 on: row 2 has no prior day, its 0 is not a return.
    last_eq_row = len(eq_curve) + 1
    ret = f"Equity_Curve!$E$3:$E${last_eq_row}"
    total_ret = f"Equity_Curve!$F${last_eq_row}"

    kpis = [
        ("NAV", '=IFERROR(LOOKUP(2,1/(Equity_Curve!$D:$D<>""),Equity_Curve!$D:$D),0)', "#,##0.00"),
        ("Total Return", '=IFERROR(LOOKUP(2,1/(Equity_Curve!$F:$F<>""),Equity_Curve!$F:$F),0)', SIGNED_PCT),
        ("Daily Return", '=IFERROR(LOOKUP(2,1/(Equity_Curve!$E:$E<>""),Equity_Curve!$E:$E),0)', SIGNED_PCT),
        ("Volatility (Ann.)", f"=IFERROR(STDEV({ret})*SQRT(252),0)", "0.00%"),
        (
            "Sharpe Ratio",
            f"=IFERROR((AVERAGE({ret})*252-RiskFreeRate)"
            f"/(STDEV({ret})*SQRT(252)),0)",
            "0.00",
        ),
        (
            # Downside deviation: target 0, sqrt(mean of squared negative returns over all days).
            "Sortino (target 0)",
            f"=IFERROR((AVERAGE({ret})*252-RiskFreeRate)"
            f"/(SQRT(SUMPRODUCT(({ret}<0)*{ret}^2)/COUNT({ret}))*SQRT(252)),0)",
            "0.00",
        ),
        ("Max Drawdown", "=IFERROR(MIN(Equity_Curve!$H:$H),0)", "0.00%"),
        (
            "Calmar Ratio",
            f"=IFERROR(((1+{total_ret})^(252/COUNT({ret}))-1)/ABS(MIN(Equity_Curve!$H:$H)),0)",
            "0.00",
        ),
        ("VaR (95% Daily)", "=IFERROR(PERCENTILE(Equity_Curve!$E:$E,0.05),0)", "0.00%"),
        ('# Positions', '=COUNTIF(Positions!$E:$E,">0")', "0"),
        # Text headers and the "" of an empty table add nothing to the SUMIF.
        ("Non-Base Currency", '=SUMIF(Currency_Exposure!$A:$A,"<>"&BaseCcy,Currency_Exposure!$C:$C)', "0.00%"),
    ]
    # Drag = commissions / starting NAV. Total Commissions lands on row 4 + len(kpis).
    kpis += [
        ("Total Commissions", commissions, "#,##0.00"),
        ("Commission Drag", f"=IFERROR(C{4 + len(kpis)}/Equity_Curve!$D$2,0)", "0.00%"),
    ]
    # Beta = slope of portfolio on benchmark daily returns; Alpha (Jensen, annualised) reads the Beta cell.
    bench_ret = f"Equity_Curve!$J$3:$J${last_eq_row}"
    beta_cell = f"C{4 + len(kpis) + 2}"
    rf_daily = "RiskFreeRate/252"
    kpis += [
        ("Benchmark", bench_ticker, "@"),
        ("Benchmark Return", f'=IF(COUNT({bench_ret})=0,"n/a",Equity_Curve!$K${last_eq_row})', SIGNED_PCT),
        ("Beta", f'=IFERROR(SLOPE({ret},{bench_ret}),"n/a")', "0.00"),
        (
            "Alpha (Ann.)",
            f'=IFERROR((AVERAGE({ret})-{rf_daily}-{beta_cell}*(AVERAGE({bench_ret})-{rf_daily}))*252,"n/a")',
            SIGNED_PCT,
        ),
        ("Start Date", "=StartDate", "@"),
    ]

    for idx, (label, formula, fmt) in enumerate(kpis):
        row = 4 + idx
        ws.cell(row=row, column=2, value=label).font = KPI_LABEL_FONT
        val_cell = ws.cell(row=row, column=3, value=formula)
        val_cell.font = KPI_VALUE_FONT
        val_cell.number_format = fmt

    for label in ("Total Return", "Daily Return"):
        cell = ws.cell(row=4 + [k[0] for k in kpis].index(label), column=3)
        ws.conditional_formatting.add(cell.coordinate, CellIsRule(operator="greaterThan", formula=["0"], font=GREEN_FONT))
        ws.conditional_formatting.add(cell.coordinate, CellIsRule(operator="lessThan", formula=["0"], font=RED_FONT))

    matrix_last_row = _write_correlation_matrix(ws, market, positions, start_col=5, start_row=4)
    nav_row = max(4 + len(kpis) - 1, matrix_last_row) + 2
    pie_row = nav_row + _rows_for(NAV_CHART_HEIGHT_CM) + 1
    _write_nav_chart(wb, ws, eq_curve, nav_row)
    _write_allocation_chart(wb, ws, positions, pie_row)
    # Pie is 14 cm high. Start in column E: B-D are too narrow/wide for a month grid.
    monthly_row = pie_row + _rows_for(14) + 1
    _write_monthly_returns(ws, eq_curve, start_col=5, start_row=monthly_row)
    # The grid takes a title row, a header row and one row per year; leave one blank row below it.
    years = len(compute_monthly_returns(eq_curve)) if not eq_curve.empty else 0
    contribution_row = monthly_row + 2 + years + 1
    _write_contribution_chart(wb, ws, contribution, anchor_row=contribution_row)
    # Same height as the contribution chart (max(8, 1 + 0.8 * n) cm), if it was drawn.
    contribution_rows = _rows_for(max(8, 1 + 0.8 * len(contribution))) if not contribution.empty else 0
    benchmark_row = contribution_row + contribution_rows + 1
    _write_benchmark_chart(wb, ws, eq_curve, anchor_row=benchmark_row)
    # Benchmark chart is 10 cm high; the drift chart (only drawn when targets exist) goes below it.
    if drift_rows:
        _write_drift_chart(wb, ws, drift_rows, anchor_row=benchmark_row + _rows_for(10) + 1)


def compute_monthly_returns(eq_curve: pd.DataFrame) -> pd.DataFrame:
    """Years x months (1..12) of NAV returns plus a compounded "Year" column."""
    nav = (eq_curve["Portfolio_Value"] + eq_curve["Cash_Balance"]).set_axis(pd.to_datetime(eq_curve["Date"]))
    month_end = nav.groupby(nav.index.to_period("M")).last()
    # Each month is measured from the previous month-end; the first month from the first NAV row.
    monthly = month_end / month_end.shift(1).fillna(nav.iloc[0]) - 1
    grid = monthly.groupby([monthly.index.year, monthly.index.month]).first().unstack()
    grid = grid.reindex(columns=range(1, 13))
    grid["Year"] = (1 + monthly).groupby(monthly.index.year).prod() - 1
    return grid


def _write_monthly_returns(ws, eq_curve: pd.DataFrame, start_col: int, start_row: int) -> None:
    if eq_curve.empty:
        return
    grid = compute_monthly_returns(eq_curve)

    ws.cell(row=start_row, column=start_col, value="Monthly Returns").font = KPI_LABEL_FONT
    ws.cell(row=start_row, column=start_col + 2, value="Year = compounded monthly returns")

    months = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec", "Year"]
    for ci, m in enumerate(months):
        cell = ws.cell(row=start_row + 1, column=start_col + 1 + ci, value=m)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")

    for ri, (year, row) in enumerate(grid.iterrows()):
        ws.cell(row=start_row + 2 + ri, column=start_col, value=int(year)).font = Font(name="Segoe UI", bold=True, size=10)
        for ci, val in enumerate(row):
            if pd.notna(val):
                cell = ws.cell(row=start_row + 2 + ri, column=start_col + 1 + ci, value=float(val))
                cell.number_format = "0.0%"

    rng = (f"{get_column_letter(start_col + 1)}{start_row + 2}:"
           f"{get_column_letter(start_col + 13)}{start_row + 1 + len(grid)}")
    # Monthly returns are small, so the scale saturates at +-10%, not the +-1 of correlations.
    ws.conditional_formatting.add(
        rng,
        ColorScaleRule(
            start_type="num", start_value=-0.1, start_color="B81D13",
            mid_type="num", mid_value=0, mid_color="FFFFFF",
            end_type="num", end_value=0.1, end_color="007A33",
        ),
    )

    for ci in range(start_col + 1, start_col + 14):
        ws.column_dimensions[get_column_letter(ci)].width = max(
            ws.column_dimensions[get_column_letter(ci)].width or 0, 9,
        )


def _write_correlation_matrix(ws, market: pd.DataFrame, positions: pd.DataFrame,
                              start_col: int, start_row: int) -> int:
    """Write the matrix and return the last row it occupies."""
    ws.cell(row=start_row - 1, column=start_col, value="Correlation Matrix").font = KPI_LABEL_FONT

    if market.empty:
        return start_row - 1

    wide = market.pivot_table(index="Date", columns="Ticker", values="Close")
    if wide.shape[1] < 2:
        return start_row - 1

    returns = wide.pct_change().dropna()
    if returns.empty:
        return start_row - 1

    corr = returns.corr()
    tickers = list(corr.columns)
    names = dict(zip(positions["Ticker"], positions["Company"])) if not positions.empty else {}

    for ci, t in enumerate(tickers):
        cell = ws.cell(row=start_row, column=start_col + 1 + ci, value=t)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")

    for ri, t_row in enumerate(tickers):
        label_cell = ws.cell(row=start_row + 1 + ri, column=start_col, value=f"{t_row}  {names[t_row]}" if names.get(t_row, t_row) != t_row else t_row)
        label_cell.font = Font(name="Segoe UI", bold=True, size=10)
        for ci, t_col in enumerate(tickers):
            val = corr.loc[t_row, t_col]
            cell = ws.cell(row=start_row + 1 + ri, column=start_col + 1 + ci, value=round(float(val), 2))
            cell.number_format = "0.00"

    data_start = f"{get_column_letter(start_col + 1)}{start_row + 1}"
    data_end = f"{get_column_letter(start_col + len(tickers))}{start_row + len(tickers)}"
    rng = f"{data_start}:{data_end}"
    ws.conditional_formatting.add(
        rng,
        ColorScaleRule(
            start_type="num", start_value=-1, start_color="B81D13",
            mid_type="num", mid_value=0, mid_color="FFFFFF",
            end_type="num", end_value=1, end_color="007A33",
        ),
    )

    ws.column_dimensions[get_column_letter(start_col)].width = max(
        [len(str(ws.cell(row=start_row + 1 + i, column=start_col).value)) for i in range(len(tickers))]
    ) + 2

    for ci in range(start_col + 1, start_col + len(tickers) + 1):
        ws.column_dimensions[get_column_letter(ci)].width = max(
            ws.column_dimensions[get_column_letter(ci)].width or 0, 10,
        )
    return start_row + len(tickers)


def _write_nav_chart(wb, ws, eq_curve: pd.DataFrame, anchor_row: int) -> None:
    if eq_curve.empty:
        return

    eq_ws = wb["Equity_Curve"]
    num_rows = len(eq_curve)

    chart = LineChart()
    chart.title = "Portfolio NAV"
    chart.style = 10
    chart.y_axis.title = "NAV ($)"
    chart.x_axis.title = "Date"
    chart.width = 32
    chart.height = NAV_CHART_HEIGHT_CM
    chart.y_axis.numFmt = "#,##0"

    data_ref = ChartRef(eq_ws, min_col=4, min_row=1, max_row=num_rows + 1)
    cats_ref = ChartRef(eq_ws, min_col=1, min_row=2, max_row=num_rows + 1)
    chart.add_data(data_ref, titles_from_data=True)
    chart.set_categories(cats_ref)
    chart.legend = None

    series = chart.series[0]
    series.graphicalProperties.line.width = 22000

    ws.add_chart(chart, f"B{anchor_row}")


def _write_allocation_chart(wb, ws, positions: pd.DataFrame, anchor_row: int) -> None:
    if positions.empty:
        return

    pos_ws = wb["Positions"]
    num_rows = len(positions)

    chart = PieChart()
    chart.title = "Portfolio Allocation"
    chart.style = 10
    chart.width = 18
    chart.height = 14

    data_ref = ChartRef(pos_ws, min_col=10, min_row=1, max_row=num_rows + 1)
    cats_ref = ChartRef(pos_ws, min_col=2, min_row=2, max_row=num_rows + 1)
    chart.add_data(data_ref, titles_from_data=True)
    chart.set_categories(cats_ref)

    chart.dataLabels = DataLabelList()
    chart.dataLabels.showPercent = True
    chart.dataLabels.showVal = False

    ws.add_chart(chart, f"B{anchor_row}")


def _write_contribution_chart(wb, ws, contribution: pd.DataFrame, anchor_row: int) -> None:
    if contribution.empty:
        return

    num_rows = len(contribution)
    chart = BarChart()
    chart.type = "bar"
    chart.title = "Contribution to Return"
    chart.style = 10
    chart.width = 18
    chart.height = max(8, 1 + 0.8 * num_rows)
    chart.legend = None
    chart.y_axis.numFmt = "0.0%"
    # Horizontal bars list the first category at the bottom; flip so the top contributor leads.
    chart.x_axis.scaling.orientation = "maxMin"
    chart.x_axis.delete = False
    chart.x_axis.tickLblPos = "low"  # keep ticker labels off the bars when values go negative
    chart.y_axis.delete = False

    c_ws = wb["Contribution"]
    chart.add_data(ChartRef(c_ws, min_col=7, min_row=1, max_row=num_rows + 1), titles_from_data=True)
    chart.set_categories(ChartRef(c_ws, min_col=1, min_row=2, max_row=num_rows + 1))
    ws.add_chart(chart, f"B{anchor_row}")


def _write_benchmark_chart(wb, ws, eq_curve: pd.DataFrame, anchor_row: int) -> None:
    if not _has_benchmark(eq_curve):
        return

    eq_ws = wb["Equity_Curve"]
    num_rows = len(eq_curve)
    chart = LineChart()
    chart.title = "Portfolio vs Benchmark (cumulative return)"
    chart.style = 10
    chart.width = 32
    chart.height = 10
    chart.legend.position = "b"
    chart.y_axis.numFmt = "0%"
    chart.x_axis.delete = False
    chart.y_axis.delete = False

    # F = Cumulative_Return, K = Benchmark_Cumulative
    for col in (6, 11):
        chart.add_data(ChartRef(eq_ws, min_col=col, min_row=1, max_row=num_rows + 1), titles_from_data=True)
    chart.set_categories(ChartRef(eq_ws, min_col=1, min_row=2, max_row=num_rows + 1))
    ws.add_chart(chart, f"B{anchor_row}")


def _write_drift_chart(wb, ws, num_rows: int, anchor_row: int) -> None:
    chart = BarChart()
    chart.type = "bar"
    chart.title = "Drift from Target"
    chart.style = 10
    chart.width = 18
    chart.height = max(8, 1 + 0.8 * num_rows)
    chart.legend = None
    chart.y_axis.numFmt = "0.0%"
    chart.x_axis.scaling.orientation = "maxMin"  # first ticker on top, as in the contribution chart
    chart.x_axis.delete = False
    chart.x_axis.tickLblPos = "low"
    chart.y_axis.delete = False

    r_ws = wb["Rebalancing"]
    chart.add_data(ChartRef(r_ws, min_col=4, min_row=1, max_row=num_rows + 1), titles_from_data=True)
    chart.set_categories(ChartRef(r_ws, min_col=1, min_row=2, max_row=num_rows + 1))
    ws.add_chart(chart, f"B{anchor_row}")


def main() -> None:
    args = parse_args()
    workbook_path = Path(args.workbook)
    if not workbook_path.exists():
        raise FileNotFoundError(f"Workbook not found: {workbook_path}")

    wb = load_workbook(workbook_path)

    start_capital = float(_read_setting(wb, "Start Capital", 100000))

    raw_trades = _read_table_df(wb, "tbl_TradeLog")
    raw_market = _read_table_df(wb, "tbl_MarketData")
    raw_reference = _read_table_df(wb, "tbl_ReferenceData")
    exposure = _read_table_df(wb, "tbl_ReferenceExposure")
    targets = _read_targets(wb)  # read before the sheet is rebuilt

    trades = _prepare_trades(raw_trades)
    market = _prepare_market(raw_market)
    base = str(_read_setting(wb, "Base Currency", "USD")).upper()
    convert_to_base(trades, market, base)
    reference = _prepare_reference(raw_reference)

    if trades.empty:
        print("Trade_Log is empty, writing empty analytics sheets")

    trading_dates = sorted(market["Date"].unique()) if not market.empty else []

    daily_units = compute_daily_units(trades, trading_dates)
    daily_cash = compute_daily_cash(trades, trading_dates, start_capital)
    positions = compute_positions(trades, market, reference)
    eq_curve = compute_equity_curve(daily_units, daily_cash, market)
    bench_ticker = str(_read_setting(wb, "Benchmark", "URTH")).strip().upper()
    if not eq_curve.empty:
        eq_curve["Benchmark_Close"] = fetch_benchmark(bench_ticker, base, list(eq_curve["Date"])).to_numpy()
    contribution = compute_contribution(trades, market, daily_units, eq_curve, reference)
    by_currency = compute_currency_exposure(positions, trades)
    commissions = compute_total_commissions(trades, eq_curve)

    _write_daily_units(wb, daily_units)
    _write_daily_cash(wb, daily_cash)
    _write_positions(wb, positions)
    _write_equity_curve(wb, eq_curve)
    _write_exposure(wb, "Sector_Exposure", "tbl_SectorExposure", "Sector", positions, exposure)
    _write_exposure(wb, "Geographic_Exposure", "tbl_GeoExposure", "Region", positions, exposure)
    if "Contribution" not in wb.sheetnames:
        wb.create_sheet("Contribution", wb.sheetnames.index("Geographic_Exposure") + 1)
    _write_contribution(wb, contribution)
    if "Currency_Exposure" not in wb.sheetnames:
        wb.create_sheet("Currency_Exposure", wb.sheetnames.index("Geographic_Exposure") + 1)
    _write_currency_exposure(wb, by_currency)
    if "Rebalancing" not in wb.sheetnames:
        wb.create_sheet("Rebalancing", wb.sheetnames.index("Currency_Exposure") + 1)
    rebalancing_rows = _write_rebalancing(wb, positions, targets)
    drift_rows = rebalancing_rows if targets else 0  # no chart without a target
    _write_dashboard(wb, eq_curve, positions, market, contribution, commissions, bench_ticker, drift_rows)

    wb.save(workbook_path)
    n_dates = len(trading_dates)
    n_tickers = len(positions) if not positions.empty else 0
    print(f"Analytics rebuilt: {n_tickers} positions, {n_dates} trading days -> {workbook_path}")


if __name__ == "__main__":
    main()
